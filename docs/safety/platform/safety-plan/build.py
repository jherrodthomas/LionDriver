#!/usr/bin/env python3
"""Build the safety plan workbook from safety-plan.json.

safety-plan.json is the source of truth; the .xlsx is generated from it and
should never be edited by hand. The base workbook comes from the safety-plan
generator (12 tabs, ISO 26262-2 §6). That generator hard-codes the Resources,
Anomaly Resolution and References tabs, so this script fills those three from
the JSON afterwards and adds notes to the confirmation reviews.

Usage:
    python build.py <path to generate_safety_plan.py>
"""
import json
import os
import subprocess
import sys

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = os.path.dirname(os.path.abspath(__file__))
SOURCE = os.path.join(HERE, "safety-plan.json")
OUTPUT = os.path.join(HERE, "LD-SPL-001_safety-plan.xlsx")

THIN = Side(border_style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BODY = Font(name="Calibri", size=10)
STRIPE = PatternFill("solid", fgColor="F2F2F2")


def write_rows(ws, first_row, rows, widths):
    """Clear everything from first_row down, then write rows with the body style."""
    for row in ws.iter_rows(min_row=first_row, max_row=max(ws.max_row, first_row)):
        for cell in row:
            cell.value = None
            cell.fill = PatternFill()
    for r, values in enumerate(rows, start=first_row):
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = BODY
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if (r - first_row) % 2:
                cell.fill = STRIPE
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def main():
    if len(sys.argv) != 2:
        print(__doc__)
        return 2
    subprocess.run([sys.executable, sys.argv[1], SOURCE, OUTPUT], check=True)
    with open(SOURCE, encoding="utf-8") as f:
        data = json.load(f)
    wb = load_workbook(OUTPUT)

    ws = wb["08_Resources"]
    write_rows(ws, 4, [(r["category"], r["description"], r["status"]) for r in data["resources"]],
               {"A": 24, "B": 90, "C": 20})

    ws = wb["09_Anomaly_Resolution"]
    ws.cell(row=3, column=1, value="Process: " + data["anomaly_process"]).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[3].height = 48
    write_rows(ws, 6, [(a["id"], a["raised_by"], a["date"], a["description"], a["severity"], a["owner"],
                        a["status"], a["resolution"]) for a in data["anomalies"]],
               {"D": 60, "H": 60})

    ws = wb["11_References"]
    write_rows(ws, 4, [(r["ref"], r["title"], r["notes"]) for r in data["references"]],
               {"A": 30, "B": 60, "C": 60})

    ws = wb["06_Confirmation_Reviews"]
    asil = {wp["wp_id"]: wp["target_asil"] for wp in data["work_products"]}
    for row in ws.iter_rows(min_row=4):
        wp_id = row[0].value
        if not wp_id:
            continue
        note = "Blocked until an independent reviewer is assigned (anomaly A01)."
        if not str(asil.get(wp_id, "")).strip().startswith(("A", "B", "C", "D")):
            note += " No ASIL: independence shown is the minimum this plan sets for non-ASIL work products."
        row[7].value = note
        row[7].alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["H"].width = 60

    wb.save(OUTPUT)
    print(f"wrote {OUTPUT}")
    return 0


if __name__ == "__main__":
    sys.exit(main())

# ruff: noqa: E501
"""Map WP-H-01 (hardware safety requirements, Markdown) onto the hw-safety-reqs-builder tab layout for review.

Requirements come from the generated trace data (trace/tools/sync_requirements.py), so they equal WP-H-01 §4.
Other tabs copy tables from WP-H-01 §6, WP-H-03, WP-H-06 and WP-A-03. Target DC is not a WP-H-01 column: it is
taken from the WP-S-03 §6 safety mechanism that lists the HWSR's parent TSR, and the column says so.
Usage: hwsr_to_xlsx.py <assurance dir> <out.xlsx>
"""
import re
import sys
from pathlib import Path

from openpyxl import Workbook

A = Path(sys.argv[1]).resolve()
OUT = Path(sys.argv[2])
sys.path.insert(0, str(A / "trace" / "tools"))
import sync_requirements as sync


def tables(path):
  out, hdr = [], None
  for line in Path(path).read_text(encoding="utf-8").splitlines():
    if line.startswith("|"):
      cells = sync.cells(line)
      if hdr is None:
        hdr = cells
        out.append((hdr, []))
      elif set("".join(cells)) <= set("-: "):
        continue
      else:
        out[-1][1].append(dict(zip(hdr, cells, strict=False)))
    else:
      hdr = None
  return out


def plain(v):
  return re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", str(v)).replace("**", "").replace("`", "")


h01 = A / "04-hardware" / "WP-H-01-hardware-safety-requirements.md"
version = re.search(r"^\| Version \| ([0-9.]+)", h01.read_text(encoding="utf-8"), re.M).group(1)
tsrs = {t["id"]: t for t in sync.tsr_items()}
hwsrs = sync.hwsr_items()

dc = {}  # TSR -> DC target of the WP-S-03 mechanism that lists it
for hdr, rows in tables(A / "03-system" / "WP-S-03-technical-safety-concept-architecture.md"):
  if hdr and hdr[0] == "SM":
    for r in rows:
      for a, b in re.findall(r"(\d{3})(?:[–-](\d{3}))?", r.get("TSR", "")):
        for n in range(int(a), int(b or a) + 1):
          dc.setdefault(f"TSR-{n:03d}", f"{plain(r['DC target'])} ({r['SM']}, WP-S-03 §6)")

hw_tsrs = [t for t in tsrs.values() if any(tag in t.get("allocated_to", "") for tag in ("P-HW", "HAR"))]
srs = []
for h in hwsrs:
  ptsr = [p for p in h["parents"] if p.startswith("TSR-")]
  srs.append([h["id"], ", ".join(ptsr) or plain(h["parent_text"]), h["asil"], plain(h["allocated_to"]), plain(h["text"]),
              next((dc[p] for p in ptsr if p in dc), "— (no WP-S-03 mechanism lists the parent TSR)"), plain(h["verification"]), h["hw_kind"], plain(h["baseline_impl"])])

fm = []
for hdr, rows in tables(A / "04-hardware" / "WP-H-03-hardware-safety-analysis-fmeda.md"):
  if hdr and hdr[0] == "FM-ID":
    for r in rows:
      fm.append([r["FM-ID"], plain(r.get("Failure mode", "")), plain(r.get("SG effect", "")), plain(r.get("Planned mechanism", "")), plain(r.get("DC planned", ""))])
hsi = []
for hdr, rows in tables(h01):
  if hdr and hdr[0] == "Proposed entry":
    for r in rows:
      hsi.append([plain(r["Proposed entry"]), plain(r["Hardware resource"]), plain(r["Software owner"]), plain(r["HWSR"])])
tests = []
for hdr, rows in tables(A / "04-hardware" / "WP-H-06-hardware-integration-verification.md"):
  if hdr and hdr[:3] == ["ID", "Title", "Verifies"]:
    for r in rows:
      tests.append([r["ID"], plain(r["Title"]), plain(r["Verifies"]), plain(r.get("Method", ""))])
indep = []
for hdr, rows in tables(A / "08-analyses" / "WP-A-03-dependent-failure-analysis.md"):
  if hdr and hdr[0] == "ID" and "Coupling factor" in hdr:
    for r in rows:
      if int(r["ID"][4:]) <= 7 or r["ID"] in ("DFI-23",):  # shared hardware resources (§4.1) and the siren path
        indep.append([r["ID"], plain(r["Initiator"])[:300], plain(r["Coupling factor"]), plain(r["Eval"]), plain(r["Required measure"])])

wb = Workbook()
wb.remove(wb.active)


def sheet(name, header, rows):
  ws = wb.create_sheet(name)
  ws.append(header)
  for r in rows:
    ws.append([plain(c) for c in r])


ws = wb.create_sheet("00_Title_Page")
for k, v in [("Title", "Hardware Safety Requirements Specification (WP-H-01), converted for review"), ("Project", "LionDriver"),
             ("Document ID", "WP-H-01"), ("Revision", version), ("Status", "Draft"), ("Author", "Assurance team (initial draft)"),
             ("Approver", "TBD (per WP-M-02)"), ("Linked TSC", "WP-S-02 v0.3 / WP-S-03 v0.3")]:
  ws.append(["", k, v])
sheet("01_Document_Control", ["Revision", "Date", "Change", "Author"],
      [("0.1", "2026-10-11", "Initial HWSRs at baseline 8b8c6ae", "Assurance team"),
       ("0.2", "2026-10-12", "HWSR-516…516b, 517, 519, 711, 711a; watchdog ≤ 0.1 s", "Assurance team")])
sheet("02_HW_TSRs_From_TSC", ["TSR ID", "Allocation", "ASIL", "Technical safety requirement"],
      [[t["id"], t["allocated_to"], t["asil"], plain(t["text"])] for t in hw_tsrs])
sheet("03_HW_Safety_Requirements", ["SR ID", "Parent TSR", "ASIL", "Allocated HW element", "Requirement text",
      "Target DC (derived from WP-S-03 §6)", "Verification method", "Kind", "Baseline"], srs)
sheet("04_Failure_Mode_Mapping", ["FM ID (WP-H-03)", "Failure mode", "SG effect", "Planned mechanism", "DC planned"], fm)
sheet("05_HSI_Reference", ["Entry", "Hardware resource", "Software owner", "HWSR"], hsi)
sheet("06_Verification_Methods", ["Method", "Meaning (WP-H-01 §4)"],
      [("A", "Analysis"), ("R", "Review"), ("T-HIL", "Bench / HIL test (WP-H-06)"), ("FI", "Fault injection"), ("T-VEH", "Vehicle test")])
sheet("07_Test_Strategy", ["Test ID", "Topic", "Verifies", "Method"], tests)
sheet("08_Independence_Considerations", ["DFI", "Initiator", "Coupling factor", "Eval", "Required measure"], indep)
wb.save(OUT)
print(f"hw_tsrs={len(hw_tsrs)} hwsrs={len(srs)} fm={len(fm)} hsi={len(hsi)} tests={len(tests)} dfi={len(indep)}")

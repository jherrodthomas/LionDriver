#!/usr/bin/env python3
# ruff: noqa: E501  (analyst data tables: one coverage entry per line)
"""LionDriver FSC completeness model (WP-C-04 §14).

The fsc-builder method expands every safety goal into a fault tree OR(input, process, output failure)
over the nodes of the system block diagram: one basic event per (safety goal, node). This file holds:

- NODES: the block diagram of the item (WP-C-01 elements split into sensor / comm / power / processor /
  software / memory / actuator / comm_out nodes, plus external elements for context);
- COVERAGE: for every (safety goal, ASIL-relevant node) basic event, the WP-C-04 FSRs that address it,
  or the reason the node cannot violate that goal.

`python3 fsc_model.py` writes the generator inputs to build/ and the coverage table
`fsc-coverage.md`. `--check` fails if the coverage table is stale, cites an FSR that WP-C-04 does
not define, leaves a basic event without FSRs or reason, has no entry at all for a basic event (an
omission is not "not relevant": NA must be written), or if a WP-C-04 FSR is not used anywhere.
`--postprocess <xlsx>` adds the coverage tab to a generated FSC workbook.
"""
from __future__ import annotations

import json
import re
import shutil
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
ASSURANCE = HERE.parent.parent
WP_C_04 = HERE.parent / "WP-C-04-functional-safety-concept.md"
GOALS = ASSURANCE / "trace" / "items" / "goals.yaml"
HARA_XLSX = HERE.parent / "hara" / "WP-C-03-hara-cartesian.xlsx"

SGS = ["SG-01", "SG-02", "SG-03", "SG-04", "SG-05", "SG-06", "SG-07", "SG-08"]

# (id, name, type, asil_relevant, inputs, outputs, redundancy_group, shared_resources, element, notes)
NODES = [
  ("N01", "Road-facing cameras", "sensor", True, [], ["N10"], None, ["device_supply", "device_pcb", "windscreen_zone"], "E-04", "Driving model input"),
  ("N02", "Driver-facing IR camera", "sensor", True, [], ["N11"], None, ["device_supply", "device_pcb", "windscreen_zone"], "E-04", "DM input"),
  ("N03", "IMU and GNSS", "sensor", True, [], ["N10"], None, ["device_supply", "device_pcb"], "E-04", "Localization, calibration"),
  ("N04", "Vehicle CAN RX: wheel speeds, steering torque sensor, brake, gas, PCM cruise", "comm_in", True, [], ["N15", "N10"], None, ["can_bus0", "harness"], "IF-02", "Gating and limiting inputs of the envelope"),
  ("N05", "EPS status 0x262", "comm_in", True, [], ["N10"], None, ["can_bus0", "harness"], "IF-02", "LKA/LTA state; read by the SoC only"),
  ("N06", "Radar tracks 0x180-0x19F", "comm_in", True, [], ["N10"], None, ["can_bus1", "harness"], "IF-02", "Lead detection"),
  ("N07", "Camera-side CAN (incl. stock PCS messages)", "comm_in", True, [], ["N15"], None, ["can_bus2", "harness"], "IF-03", "Forwarded to the car"),
  ("N08", "Vehicle 12 V supply and ignition via harness", "power_in", True, [], ["N13", "N14", "N17"], None, ["vehicle_12v", "harness"], "IF-08", "Single supply for SoC and MCU"),
  ("N09", "SoC-to-panda command stream (SPI)", "comm_in", True, ["N10"], ["N15"], None, ["device_pcb", "spi"], "IF-04", "Commands, heartbeat, safety-mode requests"),
  ("N10", "SoC driving stack (modeld, plannerd, controlsd, card, pandad)", "software", True, ["N01", "N03", "N04", "N05", "N06"], ["N09"], None, ["soc", "device_supply"], "E-01/E-02", "QM"),
  ("N11", "SoC driver monitoring (dmonitoringmodeld, dmonitoringd)", "software", True, ["N02"], ["N12"], None, ["soc", "device_supply"], "E-01/DM", "QM"),
  ("N12", "SoC state machine and alerting (selfdrived, soundd, ui)", "software", True, ["N10", "N11"], ["N19", "N09"], None, ["soc", "device_supply"], "E-01/HMI", "QM"),
  ("N13", "SoC hardware", "processor", True, ["N08"], ["N10", "N11", "N12"], None, ["soc", "device_supply", "device_pcb"], "E-04", "Controls MCU reset and boot lines (GAP-38)"),
  ("N14", "Safety MCU STM32H7 (panda)", "processor", True, ["N08"], ["N15"], None, ["device_supply", "device_pcb", "mcu_clock"], "E-03/E-04", "Single channel (GAP-11)"),
  ("N15", "Envelope software (opendbc Toyota safety mode + panda firmware)", "software", True, ["N04", "N07", "N09"], ["N17", "N18", "N20"], None, ["mcu", "mcu_clock"], "E-03", "The ASIL element"),
  ("N16", "Safety MCU memory (flash, RAM)", "memory", True, [], ["N15"], None, ["mcu"], "E-03/E-04", "Code, safety parameters, limiter state"),
  ("N17", "Harness intercept relay", "actuator", True, ["N15", "N08"], [], None, ["harness", "vehicle_12v"], "E-05", "Stock path when released"),
  ("N18", "Car-side CAN TX: 0x2E4/0x191 LKA, 0x343 ACC, 0x412 HUD, 0x1D2", "comm_out", True, ["N15"], [], None, ["can_bus0", "harness"], "IF-01", "Only path to the actuators"),
  ("N19", "Device display and speaker", "actuator", True, ["N12"], [], None, ["soc", "device_supply"], "E-04/HMI", "Take-over requests, engagement state"),
  ("N20", "Panda buzzer (siren)", "actuator", True, ["N15"], [], None, ["mcu", "device_supply"], "E-04/HMI", "SoC-independent acoustic warning"),
  ("X01", "Toyota EPS", "actuator", False, ["N18"], [], None, [], "EXT-EPS", "External; AOU-01/02"),
  ("X02", "Toyota PCM/ECM and brake actuator", "actuator", False, ["N18"], [], None, [], "EXT-PCM", "External; AOU-03/05"),
  ("X03", "Instrument cluster (cruise indicator, LKAS HUD)", "actuator", False, ["N18"], [], None, [], "EXT-CLU", "External; AOU-12"),
  ("X04", "Toyota forward camera ECU (incl. PCS)", "sensor", False, [], ["N07"], None, [], "EXT", "External"),
  ("X05", "Toyota forward radar", "sensor", False, [], ["N06"], None, [], "EXT", "External"),
]

NA = None  # marker: node cannot violate this goal; reason in NOT_RELEVANT

# COVERAGE[node][sg] = ([FSR IDs], note)  or  NA
COVERAGE: dict[str, dict[str, tuple[list[str], str] | None]] = {
  "N01": {"SG-01": (["FSR-01.01", "FSR-01.02", "FSR-01.03"], "Wrong perception is bounded by the envelope limits"),
          "SG-02": (["FSR-02.02", "FSR-02.03"], "Camera/model input loss detected and announced"),
          "SG-03": (["FSR-03.01", "FSR-03.03"], "Bounded by acceleration limit and jerk"),
          "SG-04": (["FSR-04.01", "FSR-04.02"], "Phantom braking bounded"),
          "SG-05": NA, "SG-06": (["FSR-06.01", "FSR-06.02"], "Vision lead loss detected and announced"), "SG-07": NA, "SG-08": NA},
  "N02": {"SG-01": NA, "SG-02": (["FSR-02.07"], "DM detects its own invalid input"), "SG-03": NA,
          "SG-04": (["FSR-04.05"], "DM-initiated deceleration bounded and announced (HE-04.3)"),
          "SG-05": NA, "SG-06": (["FSR-02.07"], "As SG-02"), "SG-07": NA,
          "SG-08": (["FSR-02.07", "FSR-08.01"], "Blocked or failed DM camera detected; stricter fallback; alert time counts the camera frame time")},
  "N03": {"SG-01": (["FSR-01.01", "FSR-01.02"], "Wrong localization bounded by the envelope limits"),
          "SG-02": (["FSR-02.02"], "Localization/calibration loss detected"),
          "SG-03": (["FSR-03.01"], "Bounded"), "SG-04": (["FSR-04.01", "FSR-04.02"], "Bounded"),
          "SG-05": NA, "SG-06": (["FSR-06.01"], "Input staleness detected"), "SG-07": NA, "SG-08": NA},
  "N04": {"SG-01": (["FSR-01.06", "FSR-01.05"], "RX loss/corruption/staleness of gating inputs"),
          "SG-02": (["FSR-01.06"], "Wrong gating input may revoke authority silently; detected as RX fault"),
          "SG-03": (["FSR-01.06", "FSR-03.02"], "Gas/cruise inputs gate longitudinal authority"),
          "SG-04": (["FSR-01.06", "FSR-03.02"], "No braking without granted authority (HE-04.4)"),
          "SG-05": (["FSR-01.06", "FSR-05.01", "FSR-05.02", "FSR-05.03"], "Brake, cruise and driver-torque inputs drive the release"),
          "SG-06": (["FSR-06.03", "FSR-02.12"], "PCM ACC fault reported; brake/stability fault codes block engagement"), "SG-07": NA,
          "SG-08": (["FSR-02.07"], "A wrong steering-torque or gas input can read as driver attention in wheel-touch mode (GAP-21)")},
  "N05": {"SG-01": NA, "SG-02": (["FSR-02.01", "FSR-02.12"], "EPS LKA fault states; EPS fault codes block engagement"), "SG-03": NA, "SG-04": NA, "SG-05": NA, "SG-06": NA, "SG-07": NA, "SG-08": NA},
  "N06": {"SG-01": NA, "SG-02": NA, "SG-03": (["FSR-03.01", "FSR-03.03"], "False lead release bounded"),
          "SG-04": (["FSR-04.01", "FSR-04.02"], "Ghost target braking bounded"), "SG-05": NA,
          "SG-06": (["FSR-06.01", "FSR-06.02"], "Radar fault or staleness detected and announced"), "SG-07": NA, "SG-08": NA},
  "N07": {"SG-01": (["FSR-01.11"], "Camera LKA seen on car side = relay malfunction"), "SG-02": NA, "SG-03": (["FSR-01.11"], "As SG-01 for ACC"),
          "SG-04": NA, "SG-05": NA, "SG-06": NA, "SG-07": (["FSR-07.01", "FSR-07.03"], "Forwarding and stock path"), "SG-08": NA},
  "N08": {"SG-01": (["FSR-01.09"], "Loss of supply removes actuation (fail-silent); brown-out via FSR-01.09 (TSR-505)"),
          "SG-02": (["FSR-05.05", "FSR-01.13"], "On power loss the cluster cruise indicator and the EPS timeout inform the driver; no item warning possible"),
          "SG-03": (["FSR-01.09"], "Fail-silent"), "SG-04": (["FSR-01.09", "FSR-04.03"], "Fail-silent; PCM inactive transition bounded (AOU-05R)"),
          "SG-05": (["FSR-01.09"], "Loss of supply releases control"),
          "SG-06": (["FSR-05.05"], "Cluster shows cruise off when ACC stops"), "SG-07": (["FSR-07.03"], "Relay returns to stock path on power loss"),
          "SG-08": (["FSR-01.09"], "Loss of supply ends automated control (fail-silent)")},
  "N09": {"SG-01": (["FSR-01.07", "FSR-01.08"], "Stale or corrupted commands"), "SG-02": (["FSR-01.07", "FSR-02.05"], "SoC loss: actuation removed, independent warning"),
          "SG-03": (["FSR-01.07", "FSR-01.08"], "As SG-01"), "SG-04": (["FSR-01.07", "FSR-01.08"], "As SG-01"), "SG-05": NA,
          "SG-06": (["FSR-01.07", "FSR-02.05"], "As SG-02"), "SG-07": (["FSR-01.10"], "Safety mode cannot be changed over SPI while a car mode is active"), "SG-08": NA},
  "N10": {"SG-01": (["FSR-01.01", "FSR-01.02", "FSR-01.03", "FSR-01.04", "FSR-01.14"], "QM stack; effect bounded by the envelope (T-03)"),
          "SG-02": (["FSR-02.02", "FSR-02.03", "FSR-01.07", "FSR-02.05"], "QM detection backed by envelope staleness detection"),
          "SG-03": (["FSR-03.01", "FSR-03.02", "FSR-03.03", "FSR-03.05"], "Bounded by envelope"),
          "SG-04": (["FSR-04.01", "FSR-04.02", "FSR-04.03", "FSR-04.06"], "Bounded by envelope; false FCW (HE-04.5)"),
          "SG-05": (["FSR-05.01", "FSR-05.02", "FSR-05.04"], "Release is enforced by the envelope independently of the SoC"),
          "SG-06": (["FSR-06.01", "FSR-06.02", "FSR-06.04"], "Plan/lead staleness and non-finite commands"),
          "SG-07": (["FSR-07.02"], "SoC must not originate PCS messages; enforced by the TX whitelist"),
          "SG-08": (["FSR-02.09"], "The driving stack executes the no-response reaction")},
  "N11": {"SG-01": NA, "SG-02": (["FSR-02.06", "FSR-02.07", "FSR-02.08"], "DM chain (QM, §7)"), "SG-03": NA,
          "SG-04": (["FSR-04.05", "FSR-02.09"], "DM forced deceleration (HE-04.3)"), "SG-05": NA,
          "SG-06": (["FSR-02.06", "FSR-02.07", "FSR-02.09", "FSR-06.05"], "DM chain (QM, §7); FSR-06.05 extends the DM FSRs to SG-06"), "SG-07": NA,
          "SG-08": (["FSR-02.06", "FSR-02.07", "FSR-02.08", "FSR-02.09", "FSR-08.01"], "DM chain (ASIL A for SG-08, §7): detection, self-check, no bypass, no-response reaction, alert time")},
  "N12": {"SG-01": (["FSR-01.05"], "Engagement is granted by the envelope on the PCM edge, not by the SoC state machine"),
          "SG-02": (["FSR-02.03", "FSR-02.10", "FSR-02.11"], "Take-over request; envelope/SoC mismatch; HMI path failure (H-07)"),
          "SG-03": (["FSR-01.05"], "As SG-01"), "SG-04": (["FSR-04.05"], "Alert before item-initiated deceleration"),
          "SG-05": (["FSR-05.06", "FSR-05.05"], "Cruise mismatch; cluster indication"),
          "SG-06": (["FSR-06.02", "FSR-02.10", "FSR-02.11", "FSR-05.05"], "Take-over request; mismatch; HMI failure; cluster cruise indication (H-07)"),
          "SG-07": (["FSR-07.05"], "Stock AEB/FCW events shown (QM)"),
          "SG-08": (["FSR-08.01", "FSR-02.11"], "DM alerts reach the driver in time; HMI path failure detected")},
  "N13": {"SG-01": (["FSR-01.07", "FSR-01.10"], "SoC fault: stale commands removed; SoC cannot alter the safety mode or run non-release MCU code (GAP-38, TSR-511)"),
          "SG-02": (["FSR-01.07", "FSR-02.05"], "SoC loss: actuation removed, independent warning"),
          "SG-03": (["FSR-01.07", "FSR-01.10"], "As SG-01"), "SG-04": (["FSR-01.07", "FSR-01.10"], "As SG-01"),
          "SG-05": (["FSR-05.01", "FSR-05.02", "FSR-01.10"], "Release enforced by the envelope"),
          "SG-06": (["FSR-01.07", "FSR-02.05"], "As SG-02"), "SG-07": (["FSR-01.10", "FSR-07.03"], "Mode lock; relay released on MCU reset"),
          "SG-08": (["FSR-01.07"], "A SoC fault that stops DM also stops the commands; the envelope removes actuation")},
  "N14": {sg: (["FSR-01.09"], "MCU execution faults (watchdog, clock, memory, register)") for sg in SGS} | {"SG-07": (["FSR-01.09", "FSR-07.03"], "Execution fault -> relay to stock path")},
  "N15": {"SG-01": (["FSR-01.01", "FSR-01.02", "FSR-01.03", "FSR-01.04", "FSR-01.05", "FSR-01.06", "FSR-01.07", "FSR-01.08", "FSR-01.10", "FSR-01.11", "FSR-01.12"], "The envelope functions themselves"),
          "SG-02": (["FSR-02.04", "FSR-02.05", "FSR-02.10"], "Fault ramp, independent warning; envelope blocking reported (HE-02.4)"),
          "SG-03": (["FSR-03.01", "FSR-03.02", "FSR-03.03", "FSR-03.04"], "Longitudinal limits"),
          "SG-04": (["FSR-04.01", "FSR-04.02", "FSR-04.04"], "Deceleration limits"),
          "SG-05": (["FSR-05.01", "FSR-05.02", "FSR-05.03", "FSR-05.04"], "Release on driver action"),
          "SG-06": (["FSR-01.07", "FSR-02.05"], "Envelope backstop for longitudinal loss"),
          "SG-07": (["FSR-07.01", "FSR-07.02"], "Forwarding and PCS TX block"), "SG-08": NA},
  "N16": {sg: (["FSR-01.09", "FSR-01.10"], "Memory corruption detected; safety parameters fixed") for sg in SGS},
  "N17": {"SG-01": (["FSR-01.11"], "Relay malfunction detection"), "SG-02": NA, "SG-03": (["FSR-01.11"], "As SG-01"), "SG-04": NA, "SG-05": NA, "SG-06": NA,
          "SG-07": (["FSR-07.03", "FSR-07.06"], "Released relay = stock path; stuck in intercept -> warn and refuse engagement"), "SG-08": NA},
  "N18": {"SG-01": (["FSR-01.04", "FSR-01.12", "FSR-01.13"], "Zero torque on revocation; EPS ends torque on message loss (AOU-01R)"),
          "SG-02": (["FSR-02.04"], "Bounded ramp on fault"), "SG-03": (["FSR-03.02", "FSR-03.05", "FSR-03.06"], "Inactive value and cancel"),
          "SG-04": (["FSR-04.03"], "No deceleration step on the safe-state transition"), "SG-05": (["FSR-05.01", "FSR-05.07"], "Release reaches the actuators (GAP-46, TSR-109)"),
          "SG-06": (["FSR-06.02"], "ACC cancel on loss"), "SG-07": (["FSR-07.02", "FSR-07.04", "FSR-07.07"], "No PCS message originated; PCS priority in the PCM (AOU-04R); zero acceleration request during stock AEB"), "SG-08": NA},
  "N19": {"SG-01": NA, "SG-02": (["FSR-02.03", "FSR-02.11"], "Take-over request; HMI path failure detected (H-07)"), "SG-03": NA,
          "SG-04": (["FSR-04.05"], "Alert before item-initiated deceleration"), "SG-05": (["FSR-05.06"], "Mismatch warning"),
          "SG-06": (["FSR-06.02", "FSR-02.11"], "As SG-02"), "SG-07": NA,
          "SG-08": (["FSR-02.11", "FSR-08.01"], "DM alerts shown and sounded; HMI failure detected")},
  "N20": {"SG-01": NA, "SG-02": (["FSR-02.05"], "Independent acoustic warning"), "SG-03": NA, "SG-04": NA, "SG-05": NA,
          "SG-06": (["FSR-02.05", "FSR-06.05"], "Independent acoustic warning (via FSR-06.05)"), "SG-07": NA, "SG-08": NA},
}

NOT_RELEVANT = {
  "SG-01": "Node has no path to the lateral command or its gating",
  "SG-02": "Node failure cannot remove lateral control without the driver being informed",
  "SG-03": "Node has no path to a positive acceleration request",
  "SG-04": "Node has no path to a deceleration request",
  "SG-05": "Node is not in the release path (driver brake, cancel, steering override)",
  "SG-06": "Node failure cannot remove longitudinal deceleration silently",
  "SG-07": "Node has no path to the camera-to-vehicle PCS messages",
  "SG-08": "Node has no path to driver-state detection, the driver alerts or the no-response reaction",
}


def defined_fsrs() -> set[str]:
  return set(re.findall(r"^\| (FSR-\d\d\.\d\d) \|", WP_C_04.read_text(), re.M))


def goals() -> list[dict]:
  text = GOALS.read_text()
  out = []
  for block in text.split("\n- id: ")[1:]:
    gid = block.split("\n", 1)[0].strip()
    f = {k: (m.group(1).strip().strip('"') if (m := re.search(rf"^  {k}: (.*)$", block, re.M)) else "")
         for k in ("text", "asil", "safe_state", "parents", "ftti_ms")}
    out.append({"id": gid, **f, "parents": f["parents"].strip("[]")})
  return out


def block_diagram() -> dict:
  sg_over = {}
  for g in goals():
    try:
      sg_over[g["id"]] = {"ftti_ms": int(g["ftti_ms"])}
    except ValueError:
      sg_over[g["id"]] = {}
  return {
    "fsc_metadata": {"project": "LionDriver", "doc_id": "WP-C-04 (completeness check)", "revision": "0.2", "date": "2026-10-12", "author": "Assurance team (draft)", "approver": "Safety manager", "linked_hara": "assurance/02-concept/hara/WP-C-03-hara-cartesian.xlsx (tab 13 replaced by SG-01…SG-08, see README)"},
    "block_diagram": {"nodes": [{"id": n[0], "name": n[1], "type": n[2], "asil_relevant": n[3], "inputs": n[4], "outputs": n[5],
               "redundancy_group": n[6], "shared_resources": n[7], "notes": f"{n[8]}: {n[9]}"} for n in NODES]},
    "safety_goal_overrides": sg_over,
  }


def coverage_rows():
  rows = []
  for n in NODES:
    if not n[3]:
      continue
    for sg in SGS:
      c = COVERAGE.get(n[0], {}).get(sg, NA)
      rows.append((sg, n[0], n[1], n[2], c))
  return rows


def coverage_md() -> str:
  rows = coverage_rows()
  covered = [r for r in rows if r[4]]
  lines = ["# FSC completeness check (generated)", "",
           "Generated by [`fsc_model.py`](fsc_model.py). Do not edit by hand.", "",
           f"Block diagram: {sum(1 for n in NODES if n[3])} ASIL-relevant nodes, {sum(1 for n in NODES if not n[3])} external context nodes. Basic events (safety goal × ASIL-relevant node, fsc-builder OR(input, process, output) expansion): {len(rows)}. Addressed by WP-C-04 FSRs: {len(covered)}. Not relevant (node cannot violate the goal, reason given): {len(rows) - len(covered)}. Without FSR or reason: 0 (checked).", ""]
  for sg in SGS:
    lines += [f"## {sg}", "", "| Node | Type | FSRs | How |", "|---|---|---|---|"]
    for s, nid, name, typ, c in rows:
      if s != sg:
        continue
      if c:
        lines.append(f"| {nid} {name} | {typ} | {', '.join(c[0])} | {c[1]} |")
      else:
        lines.append(f"| {nid} {name} | {typ} | — | Not relevant: {NOT_RELEVANT[sg]} |")
    lines.append("")
  used = sorted({f for r in rows if r[4] for f in r[4][0]})
  lines += ["## FSR usage", "", f"{len(used)} of {len(defined_fsrs())} FSRs in WP-C-04 are cited above. Every WP-C-04 FSR is cited at least once (checked).", ""]
  return "\n".join(lines)


def check() -> list[str]:
  errs = []
  defined = defined_fsrs()
  used = set()
  for n in NODES:
    if n[3]:
      for sg in SGS:
        if sg not in COVERAGE.get(n[0], {}):
          errs.append(f"{sg}/{n[0]}: no coverage entry (list the FSRs, or NA with the goal's not-relevant reason)")
  for sg, nid, _, _, c in coverage_rows():
    if c is None:
      continue
    if not c[0]:
      errs.append(f"{sg}/{nid}: no FSR and no reason")
    for f in c[0]:
      used.add(f)
      if f not in defined:
        errs.append(f"{sg}/{nid}: {f} is not defined in WP-C-04")
  for n in NODES:
    if n[3] and set(COVERAGE.get(n[0], {})) - set(SGS):
      errs.append(f"{n[0]}: unknown safety goal key")
  unused = sorted(defined - used)
  if unused:
    errs.append("FSRs in WP-C-04 not cited by any basic event: " + ", ".join(unused))
  if (HERE / "fsc-coverage.md").read_text() != coverage_md():
    errs.append("fsc-coverage.md is stale - run python3 fsc_model.py")
  return errs


def write_hara_sgs(out: Path) -> None:
  """Copy the HARA workbook and replace tab 13 with the consolidated goals SG-01..SG-08 (WP-C-03 §6)."""
  from openpyxl import load_workbook
  shutil.copy(HARA_XLSX, out)
  wb = load_workbook(out)
  ws = wb["13_Safety_Goals"]
  for r in range(5, ws.max_row + 1):
    for c in range(1, 9):
      ws.cell(row=r, column=c).value = None
  for i, g in enumerate(goals(), start=5):
    for c, v in enumerate([g["id"], "Item LD-SDA", g["text"], g["asil"], g["parents"], g["text"], g["safe_state"]], start=1):
      ws.cell(row=i, column=c, value=v)
  wb.save(out)


def postprocess(path: Path) -> None:
  from openpyxl import load_workbook
  from openpyxl.styles import Alignment, Font
  wb = load_workbook(path)
  ws = wb["00_Title_Page"]
  ws.cell(row=15, column=2, value="Status").font = Font(bold=True)
  ws.cell(row=15, column=3, value="Draft: completeness check of WP-C-04 v0.2; not approved (WP-C-04 header)")
  ws = wb.create_sheet("09_LionDriver_Coverage")
  ws.append(["Generated FSR rows (tab 05) are template wording per basic event. LionDriver's FSRs are in WP-C-04; this tab maps every basic event onto them (fsc_model.py)."])
  ws.append([])
  ws.append(["Safety goal", "Node", "Type", "WP-C-04 FSRs", "How / why not relevant"])
  for c in ws[3]:
    c.font = Font(bold=True)
  for sg, nid, name, typ, c in coverage_rows():
    ws.append([sg, f"{nid} {name}", typ, ", ".join(c[0]) if c else "—", c[1] if c else "Not relevant: " + NOT_RELEVANT[sg]])
  for col, w in zip("ABCDE", (10, 50, 12, 40, 80), strict=True):
    ws.column_dimensions[col].width = w
  for row in ws.iter_rows(min_row=4):
    for cell in row:
      cell.alignment = Alignment(wrap_text=True, vertical="top")
  ws = wb["06_ASIL_Allocation"]
  note = "LionDriver allocation (WP-C-04 §4, §10): the SoC nodes N10–N13 and N19 are QM; their effect on every safety goal is bounded by the envelope (N14–N16, N18) as a safety mechanism, with freedom from interference argued in WP-A-02. No ASIL decomposition is claimed. The inherited ASILs shown below are the generator default and are not the allocation."
  ws.cell(row=2, column=1, value=note)
  wb.save(path)


def main(argv):
  if len(argv) == 3 and argv[1] == "--postprocess":
    postprocess(Path(argv[2]))
    return 0
  if "--check" in argv:
    errs = check()
    for e in errs:
      print("ERROR:", e)
    print("ok" if not errs else f"{len(errs)} error(s)")
    return 1 if errs else 0
  (HERE / "build").mkdir(exist_ok=True)
  (HERE / "build" / "block_diagram.json").write_text(json.dumps(block_diagram(), indent=1) + "\n")
  write_hara_sgs(HERE / "build" / "hara_consolidated_sgs.xlsx")
  (HERE / "fsc-coverage.md").write_text(coverage_md())
  print("written: build/block_diagram.json, build/hara_consolidated_sgs.xlsx, fsc-coverage.md")
  return 0


if __name__ == "__main__":
  sys.exit(main(sys.argv))

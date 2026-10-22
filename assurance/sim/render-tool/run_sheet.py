"""Render every row of a risk-scene sheet and write the results back (render-tool-plan.md §5, piece 5).

For each row of each method sheet: row_to_scene.py makes the scene and export_scene.py exports it (the plant or the
preset's stated assumptions); Blender renders the key-event stills (../blender/hazard_scene.py), cached by render key
so rows with the same picture share one render; compose.py draws the row's own dashboard with its own card. The
Result columns get the status, a link to the dashboard, the contact, the smallest gap and TTC, and the motion source.
The dashboards go beside the sheet in <sheet>-dashboards/ as JPEG.

Blender: `blender` on the PATH, a Python with the bpy module (this one), or LIONDRIVER_BLENDER (a command that runs a
script, e.g. "blender -b -P" or "/opt/bpy/bin/python"). The purchased Corolla mesh is used when present
(assurance/sim/assets, ../blender/hazard_scene.py), otherwise the stand-in; the status says which.

  python3 assurance/sim/render-tool/run_sheet.py my-sheet.xlsx
  python3 assurance/sim/render-tool/run_sheet.py my-sheet.xlsx --no-render     # outcome only, about a second a row
  python3 assurance/sim/render-tool/run_sheet.py my-sheet.xlsx --only F-01/M03/OS-01/W02 --samples 8
"""

from __future__ import annotations

import argparse
import importlib.util
import json
import os
import shlex
import shutil
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

import row_to_scene as rs
import template

BLENDER_DIR = rs.SIM / "blender"
OUT = rs.HERE / "out"


def blender_command() -> list[str] | None:
  if os.environ.get("LIONDRIVER_BLENDER"):
    return shlex.split(os.environ["LIONDRIVER_BLENDER"])
  if shutil.which("blender"):
    return ["blender", "-b", "-P"]
  if importlib.util.find_spec("bpy"):
    return [sys.executable]
  return None


def render(data: dict, key: str, args) -> tuple[Path, str]:
  """The cached render of a picture: renders/<key>/ with stills, screen.json and render.json."""
  d = OUT / "renders" / key
  man = d / "render.json"
  if man.exists():
    return d, json.loads(man.read_text())["vehicle"]
  cmd = blender_command()
  if cmd is None:
    raise rs.RowError("no Blender: install it, pip install bpy, or set LIONDRIVER_BLENDER")
  d.mkdir(parents=True, exist_ok=True)
  # seeded by the preset, as a gallery variant is by its base scene: the same people and plants
  scene = {**data, "id": f"rt-{key}", "variant_of": data["row"]["preset"]}
  (d / "scene.json").write_text(json.dumps(scene, indent=1))
  extra = ["--stills", "--no-blend", "--samples", str(args.samples), "--res", args.res] + (
    ["--standin"] if args.standin else [])
  sep = ["--"] if cmd[0] != sys.executable else []
  p = subprocess.run(cmd + [str(BLENDER_DIR / "hazard_scene.py")] + sep + [str(d / "scene.json")] + extra,
                     capture_output=True, text=True)
  if p.returncode or not man.exists():
    (d / "blender.log").write_text(p.stdout + p.stderr)
    shutil.rmtree(d / "stills", ignore_errors=True)
    raise rs.RowError(f"Blender failed (log: {(d / 'blender.log').relative_to(rs.REPO)})")
  return d, json.loads(man.read_text())["vehicle"]


def dashboard(data: dict, rdir: Path, dest: Path) -> None:
  """The row's own dashboard from the shared render: its scene.json (card, caption) next to the render's files."""
  d = OUT / "rows" / data["id"]
  d.mkdir(parents=True, exist_ok=True)
  (d / "scene.json").write_text(json.dumps(data, indent=1))
  for f in ("render.json", "screen.json"):
    shutil.copy(rdir / f, d / f)
  link = d / "stills"
  if link.is_symlink() or link.exists():
    link.unlink() if link.is_symlink() else shutil.rmtree(link)
  link.symlink_to(os.path.relpath(rdir / "stills", d))
  p = subprocess.run([sys.executable, str(BLENDER_DIR / "compose.py"), str(d)], capture_output=True, text=True)
  if p.returncode:
    raise rs.RowError(f"compose failed: {(p.stderr or p.stdout).strip().splitlines()[-1]}")
  dest.parent.mkdir(parents=True, exist_ok=True)
  with Image.open(d / "dashboard.png") as im:          # the gallery's size and quality (gallery.py)
    im.convert("RGB").resize((1600, 900), Image.LANCZOS).save(dest, quality=75, optimize=True)


def run(sheet: Path, args) -> int:
  wb = load_workbook(sheet)
  ctx = rs.Context.load()
  out_dir = args.out_dir or sheet.with_name(f"{sheet.stem}-dashboards")
  failed = 0
  for ad in rs.load_adapters():
    if ad["method"] not in wb.sheetnames:
      continue
    ws = wb[ad["method"]]
    h = template.header(ws)
    for r, row in template.rows(ws):
      if args.only and str(row["Row ID"]) not in args.only:
        continue
      res = {k: None for k, _, _ in rs.RESULTS}
      try:
        if not str(row.get("Motion preset") or "").strip() and str(row.get("Status") or "").startswith("cannot draw"):
          raise rs.RowError(str(row["Status"])[len("cannot draw: "):])
        sc, notes = rs.row_to_scene(row, ad, ctx)
        data = rs.export(sc, ctx)
        key = rs.render_key(data)
        res.update(rs.outcome(data), **{"Render key": key})
        if args.no_render:
          res["Status"] = "; ".join(["exported, not rendered (--no-render)"] + notes)
        else:
          rdir, vehicle = render(data, key, args)
          dest = out_dir / f"{data['id']}.jpg"
          dashboard(data, rdir, dest)
          ws.cell(r, h["Dashboard"]).hyperlink = os.path.relpath(dest, sheet.parent)
          res["Dashboard"] = dest.name
          res["Status"] = "; ".join([f"drawn ({vehicle})"] + notes)
        print(f"{row['Row ID']}: {res['Status']}; contact {res['Contact']}")
      except rs.RowError as e:
        res["Status"] = f"cannot draw: {e}"
        failed += 1
        print(f"{row['Row ID']}: {res['Status']}", file=sys.stderr)
      for k, v in res.items():
        c = ws.cell(r, h[k])
        c.value = v
        if k == "Dashboard" and v is None:
          c.hyperlink = None
  wb.save(sheet)
  return 1 if failed and args.strict else 0


def main(argv: list[str]) -> int:
  ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
  ap.add_argument("sheet", type=Path)
  ap.add_argument("--only", nargs="+", help="Row IDs to run (default: every row)")
  ap.add_argument("--no-render", action="store_true", help="export and write the outcome; no Blender")
  ap.add_argument("--out-dir", type=Path, help="where the dashboards go (default: <sheet>-dashboards beside it)")
  ap.add_argument("--standin", action="store_true", help="draw every car with the stand-in, not the purchased mesh")
  ap.add_argument("--samples", type=int, default=16, help="Cycles samples (hazard_scene.py)")
  ap.add_argument("--res", default="1280x720", help="chase still resolution (hazard_scene.py)")
  ap.add_argument("--strict", action="store_true", help="exit 1 if any row cannot be drawn")
  args = ap.parse_args(argv)
  return run(args.sheet, args)


if __name__ == "__main__":
  sys.exit(main(sys.argv[1:]))

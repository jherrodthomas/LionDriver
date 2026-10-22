"""Fill rows of a risk-scene sheet from the method's own rows (render-tool-plan.md §5, piece 3).

For the ISO 26262 HARA, a row is named by its key in hara-events.csv (WP-C-03), e.g. F-01/M03/OS-01/W02. The rating
columns are copied as the HARA writes them; the environment, condition and motion preset are picked from what the
library can draw (row_to_scene.prefill), and the user may change them after. A key the library cannot draw is still
added, with the reason in its Status column. A sheet that does not exist is started from template.xlsx; rows already
in the sheet (same Row ID) are left as they are.

  python3 assurance/sim/render-tool/prefill.py my-sheet.xlsx F-01/M03/OS-01/W02 F-01/M01/OS-19/W01
  python3 assurance/sim/render-tool/prefill.py my-sheet.xlsx --method iso26262-hara F-04/M03/OS-07/W01
"""

from __future__ import annotations

import argparse
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook

import row_to_scene as rs
import template


def main(argv: list[str]) -> int:
  ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
  ap.add_argument("sheet", type=Path)
  ap.add_argument("keys", nargs="+", help="row keys of the method's own table")
  ap.add_argument("--method", default="iso26262-hara", help="adapter id (adapters/*.yaml)")
  args = ap.parse_args(argv)
  adapter = next((a for a in rs.load_adapters() if a["id"] == args.method), None)
  if adapter is None:
    print(f"no adapter {args.method}", file=sys.stderr)
    return 2
  if not args.sheet.exists():
    shutil.copy(template.TEMPLATE, args.sheet)
  wb = load_workbook(args.sheet)
  ws = wb[adapter["method"]]
  h = template.header(ws)
  have = {str(v["Row ID"]) for _, v in template.rows(ws)}
  r = max([n for n, _ in template.rows(ws)] + [template.FIRST_ROW - 1]) + 1
  ctx, src = rs.Context.load(), rs.source_rows(adapter)
  for key in args.keys:
    if key in have:
      print(f"{key}: already in the sheet")
      continue
    try:
      row = rs.prefill(adapter, key, ctx, src)
      msg = f"{key}: {row['Motion preset']} on {row['Environment']}, {row['Condition']}"
    except rs.RowError as e:
      if key not in src:
        print(f"{key}: {e}", file=sys.stderr)
        continue
      row = {"Row ID": key, **{c["column"]: src[key][c["source"]] for c in adapter["rating"]},
             "Status": f"cannot draw: {e}"}
      msg = f"{key}: added, cannot draw: {e}"
    for k, v in row.items():
      ws.cell(r, h[k], v)
    r += 1
    print(msg)
  wb.save(args.sheet)
  return 0


if __name__ == "__main__":
  sys.exit(main(sys.argv[1:]))

"""The risk-scene renderer's spreadsheet template (render-tool-plan.md §3): template.xlsx.

One sheet per method adapter (adapters/*.yaml) with the same neutral columns (Row, Scene, Look, Motion), the method's
rating columns and the result columns run_sheet.py writes back; a Lists sheet with what the library can draw (the
dropdowns read it); a Read me sheet. Row 1 names each column's kind, row 2 is the header, row 3 a hint; rows start
at row 4. Generated, never edited by hand: the dropdowns follow the library and scenes.yaml.

  python3 assurance/sim/render-tool/template.py            # write template.xlsx
  python3 assurance/sim/render-tool/template.py --check    # fail if template.xlsx's cells differ from a fresh build
  python3 assurance/sim/render-tool/template.py --out x.xlsx
"""

from __future__ import annotations

import argparse
import datetime
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import row_to_scene as rs

TEMPLATE = rs.HERE / "template.xlsx"
FIRST_ROW = 4
DATA_ROWS = 200          # rows that carry the dropdowns
KIND_FILL = {"Row": "E7E9EE", "Scene": "FBE3C8", "Look": "D9F0DF", "Motion": "D8E4FB", "Rating": "FFF2B8",
             "Result": "E9DDF7"}
STAMP = datetime.datetime(2026, 10, 22)
README = [
  "Risk-scene renderer: one row of a risk assessment in, a scene dashboard out (assurance/sim/render-tool).",
  "",
  "1. Pick the sheet of your method. Fill one row per event from row 4 down; Row ID must be unique.",
  "2. Scene and Look columns: pick from the dropdowns (the Lists sheet says what each value is).",
  "   Environment sets the road, so it can change the outcome. Condition and Traffic change the picture only.",
  "3. Motion preset: the scene whose motion the row uses, a plant run with the real panda safety code or a",
  "   kinematic motion under stated assumptions. Speed km/h is optional. There is no column for positions or paths.",
  "4. Rating columns: as your method writes them. They are shown on the dashboard; the tool never rates anything.",
  "5. Leave the Result columns empty. run_sheet.py fills them: status, dashboard, contact, gap, TTC, motion source.",
  "",
  "A HARA row can be filled for you: prefill.py <sheet.xlsx> F-01/M03/OS-01/W02 ...",
  "Every picture is synthetic (DSR-12) and earns no evidence credit. A result that disagrees with a rating is a",
  "prompt for review, not a finding.",
]


def lists(ctx: rs.Context) -> dict[str, list[tuple]]:
  lib = ctx.lib
  return {
    "Environment": [(k, " ".join(e["title"].split()), ", ".join(e["os"])) for k, e in lib["environments"].items()],
    "Condition": [(k, c["title"]) for k, c in lib["conditions"].items()],
    "Traffic": [(rs.DEFAULT, "the preset's, or the environment's")] + [(k, t.get("title", ""))
                                                                     for k, t in lib["traffic"].items()],
    "Motion preset": [(k, rs.preset_label(s), s.get("he") or (s.get("sotif") or {}).get("id", ""),
                       s["environment"]) for k, s in ctx.presets.items()],
  }


def build(ctx: rs.Context | None = None) -> Workbook:
  ctx = ctx or rs.Context.load()
  wb = Workbook()
  wb.properties.creator = "assurance/sim/render-tool/template.py"
  wb.properties.title = "Risk-scene renderer template"
  wb.properties.created = wb.properties.modified = STAMP
  ws = wb.active
  ws.title = "Read me"
  for i, line in enumerate(README, start=1):
    ws.cell(i, 1, line or None).font = Font(bold=i == 1, size=12 if i == 1 else 11)
  ws.column_dimensions["A"].width = 120

  lst = lists(ctx)
  ls = wb.create_sheet("Lists")
  heads = {"Environment": ["Environment", "What it is", "Operational situations"],
           "Condition": ["Condition", "What it is"], "Traffic": ["Traffic", "What it is"],
           "Motion preset": ["Motion preset", "Motion", "Pictures", "Its environment"]}
  ranges, col = {}, 1
  for name, rows in lst.items():
    for j, h in enumerate(heads[name]):
      c = ls.cell(1, col + j, h)
      c.font = Font(bold=True)
      ls.column_dimensions[get_column_letter(col + j)].width = 26 if j == 0 else 40
    for i, r in enumerate(rows, start=2):
      for j, v in enumerate(r):
        ls.cell(i, col + j, v)
    letter = get_column_letter(col)
    ranges[name] = f"Lists!${letter}$2:${letter}${len(rows) + 1}"
    col += len(heads[name]) + 1

  for ad in rs.load_adapters():
    sh = wb.create_sheet(ad["method"])
    cols = rs.columns(ad)
    for j, (head, kind, hint) in enumerate(cols, start=1):
      fill = PatternFill("solid", fgColor=KIND_FILL[kind])
      for r, v in ((1, kind), (2, head), (3, hint)):
        c = sh.cell(r, j, v)
        c.fill = fill
        c.alignment = Alignment(wrap_text=True, vertical="top")
      sh.cell(1, j).font = Font(bold=True, size=9, color="555555")
      sh.cell(2, j).font = Font(bold=True)
      sh.cell(3, j).font = Font(italic=True, size=9, color="555555")
      sh.column_dimensions[get_column_letter(j)].width = 34 if head in ("Status", "Dashboard", "Contact",
                                                                         "Motion preset", "Hazard") else 18
      if head in ranges:
        dv = DataValidation(type="list", formula1=f"={ranges[head]}", allow_blank=True, showErrorMessage=True,
                            errorTitle=head, error=f"Pick a {head.lower()} from the list (Lists sheet).")
        dv.add(f"{get_column_letter(j)}{FIRST_ROW}:{get_column_letter(j)}{FIRST_ROW + DATA_ROWS - 1}")
        sh.add_data_validation(dv)
    sh.row_dimensions[3].height = 42
    sh.freeze_panes = f"B{FIRST_ROW}"
  wb.move_sheet("Lists", offset=len(wb.sheetnames))
  return wb


def header(ws) -> dict[str, int]:
  """Column number of each header (row 2) of a method sheet."""
  return {c.value: c.column for c in ws[2] if c.value}


def rows(ws) -> list[tuple[int, dict]]:
  """(sheet row, {header: value}) of every row from row 4 that has a Row ID."""
  h = header(ws)
  out = []
  for r in range(FIRST_ROW, ws.max_row + 1):
    vals = {k: ws.cell(r, c).value for k, c in h.items()}
    if vals.get("Row ID") not in (None, ""):
      out.append((r, vals))
  return out


def cells(wb) -> dict:
  out = {}
  for ws in wb.worksheets:
    out[ws.title] = [[c.value for c in row] for row in ws.iter_rows()]
    out[ws.title + ":validations"] = sorted((dv.formula1, str(dv.sqref)) for dv in ws.data_validations.dataValidation)
  return out


def main(argv: list[str]) -> int:
  ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
  ap.add_argument("--check", action="store_true", help="compare template.xlsx with a fresh build")
  ap.add_argument("--out", type=Path, default=TEMPLATE)
  args = ap.parse_args(argv)
  wb = build()
  if args.check:
    if not TEMPLATE.exists() or cells(load_workbook(TEMPLATE)) != cells(wb):
      print(f"{TEMPLATE.relative_to(rs.REPO)} is stale: run python3 assurance/sim/render-tool/template.py",
            file=sys.stderr)
      return 1
    print(f"{TEMPLATE.relative_to(rs.REPO)} current")
    return 0
  wb.save(args.out)
  print(f"wrote {args.out}")
  return 0


if __name__ == "__main__":
  sys.exit(main(sys.argv[1:]))

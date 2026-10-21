#!/usr/bin/env python3
"""Export an FMEDA YAML (assurance/08-analyses/fmea/analyses/fmeda-*.yaml) to an xlsx workbook.

Usage: tools/safety/export_fmeda_xlsx.py fmeda-sg001 [OUTPUT.xlsx] [--no-recalc]
       tools/safety/export_fmeda_xlsx.py --all      write the committed review workbooks (exports/review/<name>-review.xlsx)
       tools/safety/export_fmeda_xlsx.py --check    fail if a committed review workbook differs from a fresh build

The workbook uses the tab names of fmeda-builder (automotive skills) so fmeda-checklist-reviewer can read it, but not
its generator: that generator counts non-safety-related failure rates as single-point faults and uses one formula for
residual and latent faults. Here the classification follows rating-tables.md §6.1, the same as fmea_lint.fmeda_metrics.

The worksheet and metric tabs hold live formulas, so DC or failure-rate edits recalculate in Excel. LibreOffice
(soffice) is used to store calculated values for tools that read cached values; --no-recalc skips that step.
YAML stays the source of truth; regenerate rather than editing the xlsx.

Requires PyYAML, jsonschema and openpyxl; LibreOffice for the recalculation.
"""
import argparse
import datetime
import io
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "vehicles"))
import fmea_lint
import vehicle_register

ANALYSES = Path("assurance/08-analyses/fmea/analyses")
EXPORTS = Path("assurance/08-analyses/fmea/exports")
HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
WRAP = Alignment(wrap_text=True, vertical="top")

# Worksheet columns. Order matters for the reviewer's partial-name matching: "Lambda (FIT)" must be the first header
# containing "lambda", and "Lambda S" must come before "Lambda SPF" (the reviewer's "lambda s" also matches "lambda spf").
COLS = ["HW Element", "Part", "Element FIT", "Safety related", "Row ID", "Failure mode", "Distribution %", "Lambda (FIT)", "Violates SG directly",
        "MPF potential", "Safety mechanism", "DC %", "Coverage for", "Lambda S", "Lambda SPF", "Lambda RF", "Lambda MPF_L",
        "Classification", "DFMEA ref", "Other goals", "Notes"]
C = {name: i + 1 for i, name in enumerate(COLS)}


def col(name: str) -> str:
  from openpyxl.utils import get_column_letter
  return get_column_letter(C[name])


def header_row(ws, row: int, headers: list[str]) -> None:
  for i, h in enumerate(headers, start=1):
    cell = ws.cell(row=row, column=i, value=h)
    cell.font, cell.fill, cell.alignment = HEAD, HEAD_FILL, WRAP


def kv(ws, rows: list[tuple], start: int = 3) -> None:
  for i, (k, v) in enumerate(rows, start=start):
    ws.cell(row=i, column=2, value=k).font = Font(bold=True)
    ws.cell(row=i, column=3, value=v).alignment = WRAP
  ws.column_dimensions["B"].width = 28
  ws.column_dimensions["C"].width = 110


def target_rows(doc: dict) -> list[tuple[str, str, bool]]:
  """[(goal, ASIL, shown alongside)] for the metric tabs: each target, then its also_show ASILs."""
  rows = []
  for goal, asil, also in fmea_lint.fmeda_targets(doc):
    rows.append((goal or "", asil, False))
    rows += [(goal or "", a, True) for a in also]
  return rows


def rate_summary(doc: dict) -> str:
  kinds: dict[str, int] = {}
  for c in doc["components"]:
    kinds[c.get("rate_kind", "unstated")] = kinds.get(c.get("rate_kind", "unstated"), 0) + 1
  converted = sum(1 for c in doc["components"] if c.get("stress_converted"))
  parts = [f"{n} {k}" for k, n in sorted(kinds.items())]
  return (f"{len(doc['components'])} components: {', '.join(parts)}; {converted} converted to the mission profile (IEC 61709). "
          "Handbook set approved by the lead (WP-H-03 §2.3, DQ-068); per-part source in 03_HW_Element_Catalog. "
          "Generic rates are order-of-magnitude estimates, not evidence")


def build(doc: dict, source: str) -> Workbook:
  hdr = doc["analysis"]
  targets = target_rows(doc)
  wb = Workbook()
  updated = hdr.get("updated")
  stamp = datetime.datetime.fromisoformat(str(updated)) if updated else datetime.datetime(2026, 1, 1)
  wb.properties.creator = "tools/safety/export_fmeda_xlsx.py"
  wb.properties.title = hdr["title"]
  wb.properties.created = wb.properties.modified = stamp
  ws = wb.active
  ws.title = "00_Title_Page"
  ws["A1"] = hdr["title"]
  ws["A1"].font = Font(bold=True, size=14)
  goals = "; ".join(f"{g} ASIL {a}" + (" (shown alongside)" if alongside else "") for g, a, alongside in targets)
  kv(ws, [("Project", f"LionDriver: openpilot fork, reference configuration {vehicle_register.reference_label()}"),
          ("Document ID", f"LD-FMD-{hdr['safety_goal']}"), ("Revision", hdr["revision"]), ("Status", hdr["status"].title()),
          ("Baseline", hdr["baseline"]), ("Targets (WP-H-04 §4)", goals),
          ("LD-HARA-001 goal", f"{hdr['safety_goal']} ASIL {hdr['asil_target']}" + (" (reference only; the targets above are used)"
                                                                                   if hdr.get("targets") else "")),
          ("Author", "Claude (AI-assisted draft)"), ("Approver", "Jherrod Thomas (project maintainer); not yet reviewed"),
          ("Source", f"{source} (generated by tools/safety/export_fmeda_xlsx.py; do not edit)"), ("Scope", " ".join(hdr["scope"].split()))])

  ws = wb.create_sheet("01_Document_Control")
  header_row(ws, 1, ["Revision", "Date", "Author", "Description of change"])
  history = hdr.get("history") or [{"revision": hdr["revision"], "date": hdr.get("updated", ""), "change": "(no history recorded)"}]
  for i, h in enumerate(history, start=2):
    for j, v in enumerate([h["revision"], str(h["date"]), "Claude (AI-assisted draft)", h["change"]], start=1):
      ws.cell(row=i, column=j, value=v).alignment = WRAP
  ws.column_dimensions["D"].width = 100

  ws = wb.create_sheet("02_Assumptions")
  kv(ws, [("Mission profile", hdr.get("mission_profile", "")),
          ("FIT source", rate_summary(doc)),
          ("Environmental class", "Passenger compartment, windscreen mount (sun-loaded)"),
          ("Classification rules", " ".join(["rating-tables.md §6.1 (ISO 26262-5 §8-9): direct violation without mechanism = SPF;",
                                             "with mechanism = RF for (1-DC); multiple-point latent part = λ(1-DC latent); others safe.",
                                             "Non-safety-related elements not counted. Rows classify for the file's own goal; where another",
                                             "goal differs, 06b_Per_Goal_Rows gives that goal's classification"])),
          ("Out of scope", "; ".join(hdr.get("out_of_scope", [])))], start=1)

  ws = wb.create_sheet("03_HW_Element_Catalog")
  header_row(ws, 1, ["Element ID", "Name", "Ref. des.", "FIT", "Safety related", "DFMEA element", "Rate kind", "Handbook reference",
                     "Ref. temp. (°C)", "Converted to mission profile", "Failure rate source", "Wire ends (wiring/wire-ends.tsv)"])
  for i, c in enumerate(doc["components"], start=2):
    for j, v in enumerate([c["id"], c["part"], c.get("ref_des", "block"), c["failure_rate_fit"], "Yes" if c["safety_related"] else "No",
                           c.get("dfmea_element", ""), c.get("rate_kind", ""), c.get("handbook_ref", ""), c.get("reference_temp_c", ""),
                           "Yes" if c.get("stress_converted") else "No", c["failure_rate_source"], ", ".join(c.get("wire_ends", []))],
                          start=1):
      ws.cell(row=i, column=j, value=v).alignment = WRAP
  for k, w in zip("ABCDEFGHIJKL", (12, 50, 18, 8, 10, 14, 11, 36, 10, 12, 70, 40), strict=True):
    ws.column_dimensions[k].width = w
  ws.freeze_panes = "C2"

  dc_max: dict[str, float] = {}
  for c in doc["components"]:
    for fm in c["failure_modes"]:
      for key, dck in (("safety_mechanism_spf", "dc_spf"), ("safety_mechanism_lf", "dc_lf")):
        if fm.get(key):
          dc_max[fm[key]] = max(dc_max.get(fm[key], 0.0), fm[dck])
  ws = wb.create_sheet("05_Safety_Mechanisms_From_TSC")
  header_row(ws, 1, ["Mechanism ID", "Name", "Description", "DC % (highest claimed)", "DC basis", "Requirement", "Implementation"])
  for i, m in enumerate(doc["safety_mechanisms"], start=2):
    for j, v in enumerate([m["id"], m["description"].split(":")[0][:60], m["description"], round(100 * dc_max.get(m["id"], 0.0), 1),
                           m["dc_basis"], m.get("requirement", ""), m.get("implementation", "")], start=1):
      ws.cell(row=i, column=j, value=v).alignment = WRAP
  for k, w in zip("ABCDEFG", (14, 40, 70, 12, 70, 12, 40), strict=True):
    ws.column_dimensions[k].width = w

  def split(r, d, m, dc, sr, lam):
    """λ split formulas for one row; d, m, dc, sr, lam are cell references."""
    return {"Lambda S": f'=IF({sr}="No",0,IF(AND({d}="No",{m}="No"),{lam},0))',
            "Lambda SPF": f'=IF({sr}="No",0,IF(AND({d}="Yes",{dc}=0),{lam},0))',
            "Lambda RF": f'=IF({sr}="No",0,IF(AND({d}="Yes",{dc}>0),{lam}*(1-{dc}/100),0))',
            "Lambda MPF_L": f'=IF({sr}="No",0,IF(AND({d}="No",{m}="Yes"),{lam}*(1-{dc}/100),0))'}

  def mech_dc(fm):
    if fm.get("safety_mechanism_spf") and fm.get("safety_mechanism_lf"):
      raise ValueError(f"{fm['id']}: one mechanism per row in this layout (SPF or latent)")
    mech = fm.get("safety_mechanism_spf") or fm.get("safety_mechanism_lf") or ""
    dcv = fm.get("dc_spf") if fm.get("safety_mechanism_spf") else fm.get("dc_lf", 0.0) if mech else 0.0
    return mech, round(100 * dcv, 1)

  ws = wb.create_sheet("06_FMEDA_Worksheet")
  header_row(ws, 1, COLS)
  r = 2
  row_of: dict[str, int] = {}
  for c in doc["components"]:
    for fm in c["failure_modes"]:
      mech, dcv = mech_dc(fm)
      vals = {"HW Element": c["id"], "Part": c["part"], "Element FIT": c["failure_rate_fit"], "Safety related": "Yes" if c["safety_related"] else "No",
              "Row ID": fm["id"], "Failure mode": fm["mode"],
              "Distribution %": round(100 * fm["distribution"], 4), "Violates SG directly": "Yes" if fm["violates_sg_directly"] else "No",
              "MPF potential": "Yes" if fm["mpf_potential"] else "No", "Safety mechanism": mech, "DC %": dcv,
              "Coverage for": ("single-point" if fm.get("safety_mechanism_spf") else "latent") if mech else "",
              "DFMEA ref": fm.get("dfmea_ref", ""), "Other goals": "; ".join(f"{pg['goal']}: {pg['reason']}" for pg in fm.get("per_goal", [])),
              "Notes": fm.get("notes", "")}
      lam, d, m, dc, sr = (f"{col(n)}{r}" for n in ("Lambda (FIT)", "Violates SG directly", "MPF potential", "DC %", "Safety related"))
      vals["Lambda (FIT)"] = f"={col('Element FIT')}{r}*{col('Distribution %')}{r}/100"
      vals.update(split(r, d, m, dc, sr, lam))
      # one label per row (the reviewer accepts Safe/SPF/RF/MPF_D/MPF_L): a covered direct fault is shown as RF (its residual
      # part), a multiple-point fault as MPF_D when a latent mechanism covers it, else MPF_L
      latent = f'IF({m}="Yes",IF({dc}=0,"MPF_L","MPF_D"),"Safe")'
      vals["Classification"] = f'=IF({sr}="No","Safe",IF({d}="Yes",IF({dc}=0,"SPF","RF"),{latent}))'
      for name, v in vals.items():
        ws.cell(row=r, column=C[name], value=v).alignment = WRAP
      row_of[fm["id"]] = r
      r += 1
  last = r - 1
  for name, w in (("Part", 34), ("Failure mode", 46), ("Notes", 60), ("Other goals", 40), ("Classification", 14)):
    ws.column_dimensions[col(name)].width = w
  ws.freeze_panes = "F2"

  # rows another goal classifies differently (per_goal): that goal's split next to the row's own
  ws = wb.create_sheet("06b_Per_Goal_Rows")
  pg_cols = ["Goal", "Row ID", "Lambda (FIT)", "Safety related", "Own SPF", "Own RF", "Own MPF_L", "Violates SG directly", "MPF potential",
             "Safety mechanism", "DC %", "Lambda S", "Lambda SPF", "Lambda RF", "Lambda MPF_L", "Reason"]
  P = {n: i + 1 for i, n in enumerate(pg_cols)}
  header_row(ws, 1, pg_cols)
  pr = 2
  for c in doc["components"]:
    for row in c["failure_modes"]:
      for pg in row.get("per_goal", []):
        fm = fmea_lint.fmeda_mode(row, pg["goal"])
        mech, dcv = mech_dc(fm)
        wr = row_of[row["id"]]
        ref = lambda n, wr=wr: f"='06_FMEDA_Worksheet'!{col(n)}{wr}"
        vals = {"Goal": pg["goal"], "Row ID": row["id"], "Lambda (FIT)": ref("Lambda (FIT)"), "Safety related": ref("Safety related"),
                "Own SPF": ref("Lambda SPF"), "Own RF": ref("Lambda RF"), "Own MPF_L": ref("Lambda MPF_L"),
                "Violates SG directly": "Yes" if fm["violates_sg_directly"] else "No", "MPF potential": "Yes" if fm["mpf_potential"] else "No",
                "Safety mechanism": mech, "DC %": dcv, "Reason": pg["reason"]}
        L = lambda n, pr=pr: f"{get_column_letter(P[n])}{pr}"
        vals.update(split(pr, L("Violates SG directly"), L("MPF potential"), L("DC %"), L("Safety related"), L("Lambda (FIT)")))
        for n, v in vals.items():
          ws.cell(row=pr, column=P[n], value=v).alignment = WRAP
        pr += 1
  if pr == 2:
    ws.cell(row=2, column=1, value="(none: every row classifies the same for every target goal)")
  pg_last = max(pr - 1, 2)
  ws.column_dimensions[get_column_letter(P["Reason"])].width = 70

  def total(name):
    return f"SUM('06_FMEDA_Worksheet'!{col(name)}2:{col(name)}{last})"

  def delta(goal, own, gcol):
    """Goal split minus own split over the goal's per_goal rows."""
    rng = lambda n: f"'06b_Per_Goal_Rows'!${get_column_letter(P[n])}$2:${get_column_letter(P[n])}${pg_last}"
    crit = f"'06b_Per_Goal_Rows'!$A$2:$A${pg_last}"
    return f'(SUMIF({crit},"{goal}",{rng(gcol)})-SUMIF({crit},"{goal}",{rng(own)}))'

  lam_sr = f"SUMIF('06_FMEDA_Worksheet'!{col('Safety related')}2:{col('Safety related')}{last},\"Yes\"," \
           f"'06_FMEDA_Worksheet'!{col('Lambda (FIT)')}2:{col('Lambda (FIT)')}{last})"

  def sums(goal):
    spf = f"{total('Lambda SPF')}+{delta(goal, 'Own SPF', 'Lambda SPF')}"
    rf = f"{total('Lambda RF')}+{delta(goal, 'Own RF', 'Lambda RF')}"
    mpfl = f"{total('Lambda MPF_L')}+{delta(goal, 'Own MPF_L', 'Lambda MPF_L')}"
    return spf, rf, mpfl

  ws7 = wb.create_sheet("07_SPFM_Calculation")
  header_row(ws7, 1, ["Safety goal", "ASIL", "Σ(λ_SPF + λ_RF) FIT", "Σ(λ_safety_related) FIT", "SPFM", "Target", "Status"])
  ws8 = wb.create_sheet("08_LFM_Calculation")
  header_row(ws8, 1, ["Safety goal", "ASIL", "Σ(λ_MPF_L) FIT", "Σλ − Σ(λ_SPF + λ_RF) FIT", "LFM", "Target", "Status"])
  ws9 = wb.create_sheet("09_PMHF_Calculation")
  header_row(ws9, 1, ["Safety goal", "ASIL", "λ_SPF + λ_RF (FIT)", "Target (FIT, below)", "Status", "Note"])
  for i, (goal, asil, alongside) in enumerate(targets, start=2):
    spfm_t, lfm_t, pmhf_t = fmea_lint.FMEDA_TARGETS[asil]
    spf, rf, mpfl = sums(goal)
    lab = f"{asil} (alongside)" if alongside else asil
    ws7.append([goal, lab, f"={spf}+{rf}", f"={lam_sr}", f"=1-C{i}/D{i}", spfm_t, f'=IF(E{i}>=F{i},"PASS","FAIL")'])
    ws8.append([goal, lab, f"={mpfl}", f"={lam_sr}-({spf}+{rf})", f"=IF(D{i}=0,1,1-C{i}/D{i})", lfm_t, f'=IF(E{i}>=F{i},"PASS","FAIL")'])
    ws9.append([goal, lab, f"={spf}+{rf}", pmhf_t, f'=IF(C{i}<D{i},"PASS","FAIL")',
                "Single-point and residual part only; the dual-point term (λ_MPF_L × exposure) is not included"])
  for w in (ws7, ws8, ws9):
    for k, wd in zip("ABCDEF", (12, 16, 22, 22, 10, 10), strict=True):
      w.column_dimensions[k].width = wd

  ws = wb.create_sheet("10_Metrics_Dashboard")
  ws["A1"] = f"Metrics per goal ({hdr['title']}); {'generic' if any(c.get('rate_kind', 'generic') == 'generic' for c in doc['components']) else 'handbook'} " \
             "failure rates, not evidence of compliance"
  ws["A1"].font = Font(bold=True)
  header_row(ws, 3, ["Safety goal", "ASIL", "SPFM", "SPFM target", "LFM", "LFM target", "λ_SPF + λ_RF (FIT)", "PMHF target (FIT)", "Status"])
  for i, _ in enumerate(targets, start=2):
    row = i + 2
    for j, f in enumerate([f"='07_SPFM_Calculation'!A{i}", f"='07_SPFM_Calculation'!B{i}", f"='07_SPFM_Calculation'!E{i}",
                           f"='07_SPFM_Calculation'!F{i}", f"='08_LFM_Calculation'!E{i}", f"='08_LFM_Calculation'!F{i}",
                           f"='09_PMHF_Calculation'!C{i}", f"='09_PMHF_Calculation'!D{i}",
                           f"=IF(AND('07_SPFM_Calculation'!G{i}=\"PASS\",'08_LFM_Calculation'!G{i}=\"PASS\",'09_PMHF_Calculation'!E{i}=\"PASS\"),"
                           f"\"PASS\",\"FAIL\")"], start=1):
      ws.cell(row=row, column=j, value=f)
  ws.cell(row=len(targets) + 5, column=1, value="Total FIT (safety related)").font = Font(bold=True)
  ws.cell(row=len(targets) + 5, column=3, value="='07_SPFM_Calculation'!D2")
  for k, w in zip("ABCDEFGHI", (14, 16, 10, 12, 10, 12, 18, 18, 10), strict=True):
    ws.column_dimensions[k].width = w

  ws = wb.create_sheet("11_Open_Items")
  header_row(ws, 1, ["Item", "Description", "Status"])
  items = []
  for c in doc["components"]:
    for fm in c["failure_modes"]:
      if fm["violates_sg_directly"] and not fm.get("safety_mechanism_spf"):
        items.append(f"{fm['id']} ({c['id']}): single-point fault without a safety mechanism: {fm['mode']}")
      elif fm["mpf_potential"] and not fm.get("safety_mechanism_lf"):
        items.append(f"{fm['id']} ({c['id']}): latent fault without coverage: {fm['mode']}")
      for pg in fm.get("per_goal", []):
        g = fmea_lint.fmeda_mode(fm, pg["goal"])
        if g["violates_sg_directly"] and not g.get("safety_mechanism_spf"):
          items.append(f"{fm['id']} ({c['id']}): single-point fault for {pg['goal']} without a safety mechanism: {fm['mode']}")
  generic = sum(1 for c in doc["components"] if c.get("rate_kind", "generic") == "generic")
  if generic:
    items.append(f"{generic} of {len(doc['components'])} failure rates are generic estimates: replace with handbook or supplier "
                 "values per part (WP-H-03 §2.3, OI-10)")
  if not any(c.get("stress_converted") for c in doc["components"]):
    items.append("No rate is converted to the mission profile (IEC 61709)")
  items.append("Dual-point PMHF term not computed")
  for i, t in enumerate(items, start=2):
    ws.cell(row=i, column=1, value=f"OI-{i - 1}")
    ws.cell(row=i, column=2, value=t).alignment = WRAP
    ws.cell(row=i, column=3, value="Open")
  ws.column_dimensions["B"].width = 110

  ws = wb.create_sheet("12_References")
  header_row(ws, 1, ["Reference", "Use"])
  for i, ref in enumerate(hdr.get("references", []) + [source], start=2):
    ws.cell(row=i, column=1, value=ref)
    ws.cell(row=i, column=2, value="Source of truth" if ref == source else "Basis")
  ws.column_dimensions["A"].width = 100
  return wb


def recalc(path: Path) -> bool:
  soffice = shutil.which("soffice") or shutil.which("libreoffice")
  if not soffice:
    return False
  with tempfile.TemporaryDirectory() as tmp:
    subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", tmp, str(path)],
                   check=True, capture_output=True, timeout=180)
    shutil.move(str(Path(tmp) / path.name), path)
  return True


REVIEW = EXPORTS / "review"
SOURCES = ("fmeda-sg001", "fmeda-sg009")


QUOTED_SHEET = re.compile(r"'([A-Za-z0-9_]+)'!")


def cell_values(wb) -> dict:
  """Cell contents for --check. LibreOffice's recalculation drops the quotes around plain sheet names, so they are ignored."""
  def norm(v):
    return QUOTED_SHEET.sub(r"\1!", v) if isinstance(v, str) and v.startswith("=") else v
  return {ws.title: [[norm(c.value) for c in r] for r in ws.iter_rows()] for ws in wb.worksheets}


def main(argv=None) -> int:
  ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
  ap.add_argument("name", nargs="?", help="FMEDA file stem, e.g. fmeda-sg001 (omit with --all or --check)")
  ap.add_argument("output", nargs="?", type=Path)
  ap.add_argument("--root", type=Path, default=Path("."))
  ap.add_argument("--no-recalc", action="store_true")
  ap.add_argument("--all", action="store_true", help=f"write the committed review workbooks to {REVIEW}/<name>-review.xlsx")
  ap.add_argument("--check", action="store_true", help="fail if a committed review workbook's formulas or values differ from a fresh build")
  args = ap.parse_args(argv)
  if args.all or args.check:
    stale = []
    for name in SOURCES:
      source = ANALYSES / f"{name}.yaml"
      wb = build(fmea_lint.load_yaml((args.root / source).read_text()), str(source))
      path = args.root / REVIEW / f"{name}-review.xlsx"
      if args.check:
        buf = io.BytesIO()
        wb.save(buf)
        if not path.exists() or cell_values(load_workbook(path)) != cell_values(load_workbook(io.BytesIO(buf.getvalue()))):
          stale.append(path.as_posix())
      else:
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
        if not args.no_recalc:
          recalc(path)
    if stale:
      print("export_fmeda_xlsx: stale, run tools/safety/export_fmeda_xlsx.py --all: " + ", ".join(stale))
      return 1
    print(f"export_fmeda_xlsx: {len(SOURCES)} review workbook(s) {'current' if args.check else 'written'}")
    return 0
  if not args.name:
    ap.error("name is required without --all or --check")
  source = ANALYSES / f"{args.name}.yaml"
  doc = fmea_lint.load_yaml((args.root / source).read_text())
  out = args.output or (args.root / EXPORTS / f"{args.name}.xlsx")
  out.parent.mkdir(parents=True, exist_ok=True)
  build(doc, str(source)).save(out)
  done = False if args.no_recalc else recalc(out)
  for goal, _, _ in fmea_lint.fmeda_targets(doc) or [(None, None, None)]:
    m = fmea_lint.fmeda_metrics(doc, goal)
    print(f"{out}{' ' + goal if goal else ''}: SPFM {m['spfm']:.2%}  LFM {m['lfm']:.2%}  λSPF+λRF {m['pmhf_spf_rf_fit']:.2f} FIT"
          + ("" if done else "  (formulas not recalculated)"))
  return 0


if __name__ == "__main__":
  sys.exit(main())

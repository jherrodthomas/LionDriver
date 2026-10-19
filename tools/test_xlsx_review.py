#!/usr/bin/env python3
"""Tests for xlsx_review.py: previews current, every sheet shown, rows kept, large sheets split and linked."""

import csv
import io
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
import xlsx_review as xr


class TestPreviews(unittest.TestCase):
  @classmethod
  def setUpClass(cls):
    cls.built = {wb.src: xr.build(wb) for wb in xr.WORKBOOKS}

  def test_committed_previews_current(self):
    self.assertEqual(xr.write_all(check=True), [])

  def test_every_sheet_has_a_section(self):
    for wb in xr.WORKBOOKS:
      page = self.built[wb.src]["README.md"]
      for ws in openpyxl.load_workbook(xr.ROOT / wb.src, read_only=True).worksheets:
        self.assertIn(f"\n## {ws.title}\n", page, f"{wb.src} {ws.title}")

  def test_split_keeps_every_row_once(self):
    for wb in xr.WORKBOOKS:
      for title in wb.split:
        ws = openpyxl.load_workbook(xr.ROOT / wb.src, data_only=True)[title]
        _, header, body, _ = xr.read_sheet(ws)
        rows = []
        for name, content in self.built[wb.src].items():
          if name.startswith(xr.slug(title) + "/"):
            r = list(csv.reader(io.StringIO(content)))
            self.assertEqual(r[0], header)
            self.assertLess(len(content.encode()), xr.CSV_LIMIT, name)
            rows += r[1:]
        self.assertEqual(sorted(map(tuple, rows)), sorted(map(tuple, body)), f"{wb.src} {title}")

  def test_summary_links_resolve(self):
    for wb in xr.WORKBOOKS:
      files = self.built[wb.src]
      for link in __import__("re").findall(r"\]\(([^)#]+\.csv)\)", files["README.md"]):
        self.assertIn(link, files, f"{wb.src}: {link}")

  def test_hara_counts_match_the_worksheet(self):
    page = self.built[xr.WORKBOOKS[0].src]["README.md"]
    self.assertIn("3840 rows", page)
    self.assertIn("[🟧 C (120)](12-hara-worksheet/f-01-m01.csv)", page)

  def test_cells_are_escaped_and_asil_badged(self):
    self.assertEqual(xr.md_cell("a|b\nc"), "a\\|b<br>c")
    self.assertEqual(xr.md_cell("C", "ASIL"), "🟧 C")
    self.assertEqual(xr.md_cell("C", "Function"), "C")

  def test_large_sheet_without_split_fails(self):
    with tempfile.TemporaryDirectory() as d:
      book = openpyxl.Workbook()
      ws = book.active
      ws.title = "Big"
      ws.append(["ID", "Value"])
      for i in range(xr.LARGE + 1):
        ws.append([f"R{i}", i])
      book.save(Path(d) / "big.xlsx")
      wb = xr.Workbook("big.xlsx", "out", "Big")
      with self.assertRaisesRegex(xr.PreviewError, "name the columns to split it by"):
        xr.build(wb, Path(d))
      wb.split = {"Big": ("Missing",)}
      with self.assertRaisesRegex(xr.PreviewError, "not in the header"):
        xr.build(wb, Path(d))

  def test_deterministic(self):
    wb = xr.WORKBOOKS[2]
    self.assertEqual(xr.build(wb), self.built[wb.src])


if __name__ == "__main__":
  unittest.main()

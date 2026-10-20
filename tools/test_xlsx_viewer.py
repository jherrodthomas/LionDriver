#!/usr/bin/env python3
"""Tests for xlsx_viewer.py: fingerprint current, every workbook and sheet in the page, no row lost, page sane."""

import json
import re
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
import xlsx_viewer as xv


class TestViewer(unittest.TestCase):
  @classmethod
  def setUpClass(cls):
    cls.data = xv.collect()
    cls.page = xv.build()

  def test_fingerprint_current(self):
    stored = xv.FINGERPRINT.read_text(encoding="utf-8").split()[0]
    self.assertEqual(stored, xv.fingerprint(self.page), "run tools/xlsx_viewer.py --out … and republish")

  def test_every_workbook_and_sheet_is_in_the_page(self):
    paths = xv.workbooks()
    self.assertEqual([b["src"] for b in self.data["books"]], paths)
    self.assertEqual(paths[0], xv.FIRST)
    for b in self.data["books"]:
      names = [ws.title for ws in openpyxl.load_workbook(xv.ROOT / b["src"], read_only=True).worksheets]
      self.assertEqual([s["name"] for s in b["sheets"]], names, b["src"])

  def test_hara_worksheet_rows_kept(self):
    hara = self.data["books"][0]
    ws = next(s for s in hara["sheets"] if s["name"] == "12_HARA_Worksheet")
    self.assertEqual(len(ws["rows"]), 3840)
    S = self.data["strings"]
    asil = ws["header"].index("ASIL")
    self.assertEqual(sum(S[r[asil]] == "C" for r in ws["rows"]), 110)
    self.assertIn(asil, ws["filters"])
    self.assertEqual(ws["narrow"][:2], [ws["header"].index("HARA ID"), asil])

  def test_page_embeds_parseable_data(self):
    m = re.search(r'<script id="data" type="application/json">(.*?)</script>', self.page, re.S)
    self.assertEqual(json.loads(m.group(1)), self.data)
    self.assertLess(len(self.page.encode()), 16_000_000)
    self.assertIn("<title>Assurance Workbook Viewer</title>", self.page[:8000])

  def test_read_sheet_title_header_and_note(self):
    with tempfile.TemporaryDirectory() as d:
      book = openpyxl.Workbook()
      ws = book.active
      ws.append(["A sheet title"])
      ws.append([])
      ws.append(["ID", "Value", "Note"])
      ws.append(["R1", 1, "x|y"])
      ws.append([])
      ws.append(["A long note under the table that runs well past sixty characters in one cell"])
      book.save(Path(d) / "t.xlsx")
      above, header, body, below = xv.read_sheet(openpyxl.load_workbook(Path(d) / "t.xlsx").active)
    self.assertEqual(above, ["A sheet title"])
    self.assertEqual(header, ["ID", "Value", "Note"])
    self.assertEqual(body, [["R1", "1", "x|y"]])
    self.assertEqual(len(below), 1)

  def test_titles_and_groups(self):
    self.assertEqual(xv.title("assurance/02-concept/reviews/RR-2026-003-hara-checklist.xlsx"), "RR-2026-003 HARA checklist")
    self.assertEqual(xv.group("assurance/02-concept/reviews/RR-2026-003-hara-checklist.xlsx"), "02 Concept")

  def test_deterministic(self):
    self.assertEqual(xv.build(), self.page)


if __name__ == "__main__":
  unittest.main()
